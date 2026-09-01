"""C 智能测评题目批量导入服务。"""

from __future__ import annotations

import csv
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
import re
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.dao.assessment import QuestionBankDAO
from app.models.assessment import AssessmentQuestion
from app.schemas.assessment import QuestionCreate
from app.services.assessment_service import AssessmentService


class AssessmentImportService:
    MAX_ROWS = 1000
    HEADER_ALIASES = {
        "type": {"题型", "type"},
        "content": {"题干", "题目", "content"},
        "options": {"选项", "options"},
        "answer": {"标准答案", "答案", "answer"},
        "dimension": {"能力维度", "维度", "dimension"},
        "difficulty": {"难度", "difficulty"},
        "score": {"分值", "score"},
        "status": {"状态", "status"},
    }
    REQUIRED_HEADERS = {"type", "content", "options", "answer", "dimension", "difficulty", "score"}
    TYPE_ALIASES = {
        "single": "single",
        "单选": "single",
        "单选题": "single",
        "multi": "multi",
        "多选": "multi",
        "多选题": "multi",
        "judge": "judge",
        "判断": "judge",
        "判断题": "judge",
    }

    @classmethod
    def import_questions(
        cls, db: Session, *, bank_id: int, filename: str, content: bytes
    ) -> dict[str, Any]:
        if not QuestionBankDAO.get(db, bank_id):
            raise LookupError("题库不存在")
        rows = cls._read_rows(filename, content)
        if not rows:
            raise ValueError("导入文件没有有效数据行")
        if len(rows) > cls.MAX_ROWS:
            raise ValueError(f"单次最多导入 {cls.MAX_ROWS} 道题目")

        existing_contents = set(db.scalars(select(AssessmentQuestion.content).where(
            AssessmentQuestion.bank_id == bank_id
        )).all())
        seen_contents: set[str] = set()
        prepared: list[tuple[QuestionCreate, int]] = []
        errors: list[dict[str, Any]] = []

        for row_number, row in rows:
            try:
                question, status = cls._validate_row(bank_id, row)
                if question.content in existing_contents:
                    raise ValueError("该题干已存在于当前题库")
                if question.content in seen_contents:
                    raise ValueError("导入文件中存在重复题干")
                seen_contents.add(question.content)
                prepared.append((question, status))
            except (TypeError, ValueError) as exc:
                errors.append({
                    "row": row_number,
                    "content": cls._text(row.get("content")) or None,
                    "message": str(exc),
                })

        if errors:
            return {
                "total_count": len(rows),
                "imported_count": 0,
                "failed_count": len(errors),
                "errors": errors,
            }

        for question_body, status in prepared:
            question = AssessmentService.create_question(db, question_body)
            question.status = status
        db.flush()
        return {
            "total_count": len(rows),
            "imported_count": len(prepared),
            "failed_count": 0,
            "errors": [],
        }

    @classmethod
    def build_template(cls) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise RuntimeError("生成 Excel 模板需要安装 openpyxl") from exc

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "题目导入"
        headers = ["题型", "题干", "选项", "标准答案", "能力维度", "难度", "分值", "状态"]
        sheet.append(headers)
        sheet.append(["单选", "项目启动阶段最先应明确什么？", "项目目标|项目范围|项目预算", "项目目标", "项目管理", 1, 10, "启用"])
        sheet.append(["多选", "哪些属于风险管理动作？", "识别风险|评估影响|制定应对|忽略风险", "识别风险|评估影响|制定应对", "风险管理", 2, 10, "启用"])
        sheet.append(["判断", "服务端应负责最终判分。", "正确|错误", "正确", "数字化素养", 1, 10, "启用"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F6FED")
        widths = [12, 42, 48, 32, 18, 10, 10, 10]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        sheet.freeze_panes = "A2"
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    @classmethod
    def _read_rows(cls, filename: str, content: bytes) -> list[tuple[int, dict[str, Any]]]:
        extension = Path(filename or "").suffix.lower()
        if extension == ".xlsx":
            return cls._read_xlsx(content)
        if extension == ".csv":
            return cls._read_csv(content)
        raise ValueError("仅支持 .xlsx 或 .csv 文件")

    @classmethod
    def _read_xlsx(cls, content: bytes) -> list[tuple[int, dict[str, Any]]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("导入 Excel 需要安装 openpyxl") from exc
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            header_map = cls._header_map(headers)
            rows = []
            for row_number, values_row in enumerate(values, start=2):
                row = {key: values_row[index] if index < len(values_row) else None for key, index in header_map.items()}
                if any(cls._text(value) for value in row.values()):
                    rows.append((row_number, row))
            return rows
        except (OSError, ValueError, TypeError) as exc:
            raise ValueError(f"Excel 文件无法读取：{exc}") from exc

    @classmethod
    def _read_csv(cls, content: bytes) -> list[tuple[int, dict[str, Any]]]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV 文件必须使用 UTF-8 编码") from exc
        reader = csv.reader(StringIO(text))
        headers = next(reader, None)
        header_map = cls._header_map(headers)
        rows = []
        for row_number, values_row in enumerate(reader, start=2):
            row = {key: values_row[index] if index < len(values_row) else None for key, index in header_map.items()}
            if any(cls._text(value) for value in row.values()):
                rows.append((row_number, row))
        return rows

    @classmethod
    def _header_map(cls, headers) -> dict[str, int]:
        if not headers:
            raise ValueError("导入文件缺少表头")
        normalized = {cls._text(value).lower(): index for index, value in enumerate(headers)}
        result = {}
        for key, aliases in cls.HEADER_ALIASES.items():
            for alias in aliases:
                if alias.lower() in normalized:
                    result[key] = normalized[alias.lower()]
                    break
        missing = cls.REQUIRED_HEADERS - result.keys()
        if missing:
            labels = [next(iter(cls.HEADER_ALIASES[key])) for key in sorted(missing)]
            raise ValueError(f"缺少必要列：{'、'.join(labels)}")
        return result

    @classmethod
    def _validate_row(cls, bank_id: int, row: dict[str, Any]) -> tuple[QuestionCreate, int]:
        question_type = cls.TYPE_ALIASES.get(cls._text(row.get("type")).lower())
        if not question_type:
            raise ValueError("题型必须是单选、多选或判断")
        content = cls._text(row.get("content"))
        if not content:
            raise ValueError("题干不能为空")
        dimension = cls._text(row.get("dimension"))
        if not dimension:
            raise ValueError("能力维度不能为空")
        try:
            difficulty = int(row.get("difficulty"))
        except (TypeError, ValueError) as exc:
            raise ValueError("难度必须是 1-5 的整数") from exc
        if difficulty not in range(1, 6):
            raise ValueError("难度必须是 1-5 的整数")
        try:
            score = Decimal(cls._text(row.get("score")))
        except (InvalidOperation, ValueError) as exc:
            raise ValueError("分值必须是大于 0 的数字") from exc
        if score <= 0:
            raise ValueError("分值必须是大于 0 的数字")

        options = cls._split(row.get("options"), r"[|\n]")
        answers = cls._split(row.get("answer"), r"[|,，;；\n]")
        if question_type == "judge" and not options:
            options = ["正确", "错误"]
        if len(options) < 2:
            raise ValueError("选项至少需要 2 个")
        if len(options) != len(set(options)):
            raise ValueError("选项不能重复")
        if not answers:
            raise ValueError("标准答案不能为空")
        if question_type in {"single", "judge"} and len(answers) != 1:
            raise ValueError("单选题和判断题只能有一个标准答案")
        if any(answer not in options for answer in answers):
            raise ValueError("标准答案必须与选项文本完全一致")

        status_text = cls._text(row.get("status")).lower()
        if status_text in {"", "1", "启用", "正常", "true"}:
            status = 1
        elif status_text in {"0", "停用", "禁用", "false"}:
            status = 0
        else:
            raise ValueError("状态必须是启用或停用")
        answer: list[str] | str = answers if question_type in {"single", "multi"} else answers[0]
        try:
            body = QuestionCreate(
                bank_id=bank_id,
                type=question_type,
                content=content,
                options=options,
                answer=answer,
                dimension=dimension,
                difficulty=difficulty,
                score=score,
            )
        except ValueError as exc:
            raise ValueError(str(exc)) from exc
        return body, status

    @staticmethod
    def _split(value: Any, pattern: str) -> list[str]:
        return [item.strip() for item in re.split(pattern, str(value or "")) if item.strip()]

    @staticmethod
    def _text(value: Any) -> str:
        return str(value).strip() if value is not None else ""
