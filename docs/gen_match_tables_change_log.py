# -*- coding: utf-8 -*-
"""生成岗位匹配域建表改动记录表 Excel。"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============ 数据 ============
# 表级概览
overview_rows = [
    ["表名", "表中文名", "操作", "对应需求", "对应任务", "外键", "备注"],
    ["pos_position", "业务岗位", "新增", "M-1 岗位管理", "T-P4-01", "sys_position_id→sys_position.id；dept_id→sys_dept.id", "岗位说明书为画像化输入源；code 唯一"],
    ["match_rule", "匹配规则", "新增", "M-3 双向匹配", "T-P4-02", "—", "rule_json 存维度权重 JSON"],
    ["match_result", "匹配结果", "新增", "M-3/M-4", "T-P4-02", "position_id→pos_position.id", "talent_id 未加FK（tal_talent未建）；UNIQUE(talent_id,position_id)"],
    ["match_push_log", "预警推送日志", "新增", "M-5 储备/空缺预警", "T-P4-02", "match_id→match_result.id；message_id→msg_center.id", "type: reserve/vacancy"],
]

# 字段明细
detail_rows = [
    ["表名", "字段名", "类型", "可空", "默认值", "说明"],
    # pos_position
    ["pos_position", "id", "int", "否", "自增", "主键"],
    ["pos_position", "sys_position_id", "int", "是", "NULL", "对应基础岗位/职级（FK→sys_position.id）"],
    ["pos_position", "name", "varchar(64)", "否", "—", "岗位名称（索引）"],
    ["pos_position", "code", "varchar(64)", "否", "—", "岗位编码（唯一）"],
    ["pos_position", "dept_id", "int", "是", "NULL", "所属部门（FK→sys_dept.id）"],
    ["pos_position", "headcount", "int", "否", "0", "编制人数"],
    ["pos_position", "filled", "int", "否", "0", "已到岗人数"],
    ["pos_position", "status", "int", "否", "1", "1启用 0停用"],
    ["pos_position", "description", "text", "是", "NULL", "岗位说明书（画像化输入源）"],
    ["pos_position", "created_at", "datetime", "否", "now", "创建时间"],
    ["pos_position", "updated_at", "datetime", "否", "now", "更新时间（自动更新）"],
    # match_rule
    ["match_rule", "id", "int", "否", "自增", "主键"],
    ["match_rule", "name", "varchar(64)", "否", "—", "规则名称"],
    ["match_rule", "rule_json", "text", "是", "NULL", "维度权重JSON，如 {\"skill\":0.4,\"degree\":0.2,\"years\":0.2,\"quality\":0.2}"],
    ["match_rule", "status", "int", "否", "1", "1启用 0停用"],
    ["match_rule", "created_at", "datetime", "否", "now", "创建时间"],
    # match_result
    ["match_result", "id", "int", "否", "自增", "主键"],
    ["match_result", "talent_id", "int", "否", "—", "人才ID（索引；未加FK，tal_talent表未建）"],
    ["match_result", "position_id", "int", "否", "—", "岗位ID（FK→pos_position.id，索引）"],
    ["match_result", "score", "decimal(5,2)", "否", "0", "匹配度 0-100"],
    ["match_result", "dimension_json", "text", "是", "NULL", "各维度得分 JSON"],
    ["match_result", "explain", "text", "是", "NULL", "解释依据（Agent③生成）"],
    ["match_result", "rank", "int", "是", "NULL", "排序名次"],
    ["match_result", "status", "int", "否", "0", "0候选 1推荐 2录用"],
    ["match_result", "created_at", "datetime", "否", "now", "创建时间"],
    # match_push_log
    ["match_push_log", "id", "int", "否", "自增", "主键"],
    ["match_push_log", "match_id", "int", "否", "—", "匹配结果ID（FK→match_result.id，索引）"],
    ["match_push_log", "type", "varchar(16)", "否", "—", "reserve储备 / vacancy空缺"],
    ["match_push_log", "target_user", "varchar(64)", "是", "NULL", "目标用户（用户名/ID）"],
    ["match_push_log", "message_id", "int", "是", "NULL", "关联消息（FK→msg_center.id）"],
    ["match_push_log", "created_at", "datetime", "否", "now", "创建时间"],
]

# ============ 样式 ============
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="4F81BD")
alt_fill = PatternFill("solid", fgColor="DCE6F1")
thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_sheet(ws, headers_count, widths):
    for c in range(1, headers_count + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    for r in range(2, ws.max_row + 1):
        for c in range(1, headers_count + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = center if c in (1,) else left
            if r % 2 == 0:
                cell.fill = alt_fill


wb = Workbook()

# Sheet1 表级概览
ws1 = wb.active
ws1.title = "改动概览"
ws1.append(["数据库改动记录 — 岗位匹配域建表（2026-08-31）"])
ws1.merge_cells("A1:G1")
ws1["A1"].font = Font(bold=True, size=14)
ws1["A1"].alignment = center
ws1.append(["数据库", "云服务器 MySQL（120.77.177.232:3306/adtp_db）"])
ws1.append(["操作类型", "仅新增 4 张表，未修改/删除任何已有表"])
ws1.append(["迁移前表数", "13（12 业务表 + alembic_version）"])
ws1.append(["迁移后表数", "17"])
ws1.append(["Alembic 版本", "40ffbb1d3136 → 513968600ebc"])
ws1.append([])
ws1.append(overview_rows[0])
for row in overview_rows[1:]:
    ws1.append(row)

ws1.merge_cells("A2:B2")
ws1.merge_cells("A3:B3")
ws1.merge_cells("A4:B4")
ws1.merge_cells("A5:B5")
ws1.merge_cells("A6:B6")
style_sheet(ws1, 7, [18, 14, 10, 16, 12, 44, 46])
for r in range(2, 7):
    ws1.cell(row=r, column=1).font = Font(bold=True)

# Sheet2 字段明细
ws2 = wb.create_sheet("字段明细")
ws2.append(detail_rows[0])
for row in detail_rows[1:]:
    ws2.append(row)
style_sheet(ws2, 6, [18, 16, 14, 8, 10, 70])
# 表名分组着色（按第一列变化区分）
table_fill = {"pos_position": "E2EFDA", "match_rule": "FFF2CC", "match_result": "FCE4D6", "match_push_log": "D9E2F3"}
prev = None
for r in range(2, ws2.max_row + 1):
    t = ws2.cell(row=r, column=1).value
    if t != prev:
        fill = PatternFill("solid", fgColor=table_fill.get(t, "FFFFFF"))
        prev = t
    ws2.cell(row=r, column=1).fill = fill
    ws2.cell(row=r, column=1).font = Font(bold=True)

out = "D:/project/gitee_tanle/ai_talent/docs/岗位匹配域建表改动记录_2026-08-31.xlsx"
wb.save(out)
print("saved:", out)
